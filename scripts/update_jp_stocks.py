"""JPX 東証上場銘柄一覧(data_j.xlsx) → data/jp-stocks.json

일본 종목의 일본어 이름·검색용. 시세는 Yahoo(.T)에서 받음.
GitHub Actions(update-jp-stocks.yml)가 매월 실행. 로컬: python scripts/update_jp_stocks.py data/jp-stocks.json [xlsx경로]
"""
import datetime
import io
import json
import os
import sys
import unicodedata
import urllib.request

from openpyxl import load_workbook

URL = 'https://www.jpx.co.jp/markets/statistics-equities/misc/tvdivq0000001vg2-att/data_j.xlsx'

# 市場・商品区分 → 앱에서 쓰는 짧은 구분. 목록에 없는 구분(PRO Market, 出資証券 등)은 제외
SEGMENTS = [
    ('プライム', 'Prime'),
    ('スタンダード', 'Standard'),
    ('グロース', 'Growth'),
    ('ETF', 'ETF'),
    ('REIT', 'REIT'),
]
MIN_ITEMS = 3000  # 이보다 적으면 파일 형식이 바뀐 것으로 보고 중단


def segment_of(raw):
    raw = str(raw or '')
    for prefix, short in SEGMENTS:
        if raw.startswith(prefix):
            return short
    return None


def as_date(v):
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    if s.isdigit() and len(s) == 8:
        return f'{s[:4]}-{s[4:6]}-{s[6:]}'
    raise ValueError(f'unexpected 日付 value: {v!r}')


def as_code(v):
    # 숫자 코드는 int로 읽힐 수 있음. 2024년부터 영문 포함 코드(130A 등)도 있음
    s = str(int(v)) if isinstance(v, float) and v.is_integer() else str(v).strip()
    return s.upper()


def main(out_path, source=None):
    if source:
        data = open(source, 'rb').read()
    else:
        req = urllib.request.Request(URL, headers={'User-Agent': 'Mozilla/5.0 stock-peek-updater'})
        data = urllib.request.urlopen(req, timeout=60).read()

    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else '' for h in next(rows)]
    try:
        i_date, i_code, i_name, i_seg = (header.index(c) for c in ('日付', 'コード', '銘柄名', '市場・商品区分'))
    except ValueError as e:
        raise SystemExit(f'column not found ({e}); header={header}')

    as_of = None
    items = []
    for r in rows:
        if r is None or r[i_code] is None:
            continue
        seg = segment_of(r[i_seg])
        if not seg:
            continue
        as_of = as_of or as_date(r[i_date])
        # 전각 영숫자·반각 가나 정리 (ｉＦｒｅｅ → iFree): 표시와 검색을 단순하게
        name = unicodedata.normalize('NFKC', str(r[i_name])).strip()
        items.append([as_code(r[i_code]), name, seg])

    if len(items) < MIN_ITEMS:
        raise SystemExit(f'only {len(items)} items parsed; refusing to write')
    items.sort(key=lambda x: x[0])

    # 한 줄에 한 종목: 월간 변경분이 diff로 보이게
    lines = ',\n'.join('    ' + json.dumps(it, ensure_ascii=False) for it in items)
    os.makedirs(os.path.dirname(out_path) or '.', exist_ok=True)
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write('{\n')
        f.write(f'  "source": "JPX 東証上場銘柄一覧",\n  "asOf": "{as_of}",\n  "items": [\n{lines}\n  ]\n}}\n')
    print(f'{len(items)} items as of {as_of} -> {out_path}')


if __name__ == '__main__':
    try:
        main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
    except BaseException as e:  # noqa: BLE001 — Actions 주석으로 원인을 남김
        if os.environ.get('GITHUB_ACTIONS'):
            print(f'::error::{type(e).__name__}: {e}')
        raise
